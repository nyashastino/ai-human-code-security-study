"""Recalculate study results from the archived scan records (no server required)."""
from pathlib import Path
import csv
import json
import re
from collections import Counter
import numpy as np
from scipy import stats
import openpyxl
from openpyxl.styles import Font, Alignment
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
def read_csv(path):
    with path.open() as f:
        return list(csv.DictReader(f))
def load(path):
    return json.loads(path.read_text())
def write_csv(name, rows):
    with (ROOT/'results'/name).open('w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader(); w.writerows(rows)

sources = read_csv(ROOT/'data/human-sources.csv')
tasks = {int(r['task_id']): r['task'] for r in sources}
reviews = {r['finding_id']: r for r in read_csv(ROOT/'data/manual-review.csv')}
per_sample, vulnerabilities, hotspots, deduplicated = [], [], [], []
for group in ('AI', 'Human'):
    folder = ROOT/'scan/raw'/group.lower()
    issues = [r for r in load(folder/'issues.json') if r['type'] == 'VULNERABILITY']
    spots = load(folder/'hotspots.json')
    files = {r['key']: r['path'] for r in load(folder/'files.json')}
    metrics = load(folder/'file_measures.json')
    issue_counts, hotspot_counts, file_counts, code_lines = Counter(), Counter(), Counter(), Counter()
    groups = {}
    for path in files.values():
        file_counts[int(re.search(r'sample_(\d+)', path)[1])] += 1
    for record in metrics:
        n = int(re.search(r'sample_(\d+)', record['path'])[1])
        code_lines[n] += sum(float(m['value']) for m in record.get('measures', []) if m['metric'] == 'ncloc' and 'value' in m)
    for r in issues:
        path = files[r['component']]; n = int(re.search(r'sample_(\d+)', path)[1])
        issue_counts[n] += 1
        line = r['line']; rule = r['rule']
        if rule in ('python:S2068', 'python:S6779', 'javascript:S6437'):
            category = 'Hardcoded secret'
        elif rule == 'python:S8392':
            category = 'All-interface binding'
        elif rule == 'python:S2115':
            category = 'Database password protection'
        else:
            raise ValueError('Unmapped vulnerability rule: '+rule)
        if rule == 'javascript:S6437':
            origins = [loc for flow in r['flows'] for loc in flow['locations'] if loc.get('msg') == 'Hardcoded value assigned here']
            assert len(origins) == 1
            line = origins[0]['textRange']['startLine']
        groups.setdefault((n, path, line, category), []).append(r)
        vulnerabilities.append(dict(group=group, sample=n, file=path, line=r['line'], rule=rule, severity=r['severity'], message=r['message'], finding_id=r['key'], manual_decision=reviews[r['key']]['decision']))
    unique_counts = Counter()
    for (n, path, line, category), records in sorted(groups.items()):
        unique_counts[n] += 1
        decisions = {reviews[r['key']]['decision'] for r in records}
        assert len(decisions) == 1
        deduplicated.append(dict(group=group, sample=n, file=path, source_line=line, category=category, raw_findings=len(records), rules=';'.join(sorted({r['rule'] for r in records})), finding_ids=';'.join(r['key'] for r in records), manual_decision=decisions.pop()))
    for r in spots:
        path = files[r['component']]; n = int(re.search(r'sample_(\d+)', path)[1])
        hotspot_counts[n] += 1
        hotspots.append(dict(group=group, sample=n, file=path, line=r.get('line', ''), rule=r['ruleKey'], priority=r['vulnerabilityProbability'], message=r['message'], finding_id=r['key'], manual_decision=reviews[r['key']]['decision']))
    for n in range(1,31):
        per_sample.append(dict(group=group, sample=n, task=tasks[n], source_files=file_counts[n], code_lines=code_lines[n], raw_vulnerabilities=issue_counts[n], deduplicated_issues=unique_counts[n], hotspots=hotspot_counts[n], debug_hotspots=sum(x['group']==group and x['sample']==n and x['rule']=='python:S4507' for x in hotspots)))

write_csv('per-implementation.csv', per_sample)
write_csv('vulnerabilities.csv', vulnerabilities)
write_csv('hotspots.csv', hotspots)
write_csv('deduplicated-issues.csv', deduplicated)
a = np.array([r['raw_vulnerabilities'] for r in per_sample if r['group']=='AI'])
b = np.array([r['raw_vulnerabilities'] for r in per_sample if r['group']=='Human'])
u = stats.mannwhitneyu(a,b,alternative='two-sided',method='asymptotic',use_continuity=True)
rng = np.random.default_rng(20261011)
ab = rng.choice(a,(20000,30)); bb = rng.choice(b,(20000,30))
statistics = dict(ai_findings=int(a.sum()), human_findings=int(b.sum()), ai_hotspots=60, human_hotspots=56,
    u_ai=float(u.statistic), u_min=float(min(u.statistic,900-u.statistic)), p_two_sided=float(u.pvalue),
    mean_difference=float(a.mean()-b.mean()), mean_difference_ci95=np.quantile(ab.mean(1)-bb.mean(1),[.025,.975]).tolist(),
    cliffs_delta=float(np.sign(a[:,None]-b).mean()), cliffs_delta_ci95=np.quantile(np.sign(ab[:,:,None]-bb[:,None,:]).mean((1,2)),[.025,.975]).tolist(),
    bootstrap_replicates=20000, bootstrap_seed=20261011,
    test='Two-sided asymptotic Mann–Whitney U, with tie and continuity corrections; raw finding counts',
    deduplication='Same underlying secret assignment merged across rules/usages; other locations remain distinct')
expected = load(ROOT/'scan/reference-statistics.json')
for key in ('ai_findings','human_findings','ai_hotspots','human_hotspots','u_ai','u_min','p_two_sided','mean_difference','mean_difference_ci95','cliffs_delta','cliffs_delta_ci95'):
    assert np.allclose(statistics[key], expected[key], rtol=1e-12, atol=1e-12), key
assert len(vulnerabilities)==25 and len(hotspots)==116 and len(deduplicated)==18
assert sum(r['source_files'] for r in per_sample if r['group']=='Human')==225
assert Counter(r['group'] for r in deduplicated)=={'AI':11,'Human':7}
assert all(r['severity']=='BLOCKER' for r in vulnerabilities)
(ROOT/'results/statistics.json').write_text(json.dumps(statistics,indent=2)+'\n')

summary = [
    {'Measure':'Implementations','AI':30,'Human':30},
    {'Measure':'Scanned source files','AI':30,'Human':225},
    {'Measure':'Raw Vulnerability findings (all BLOCKER)','AI':16,'Human':9},
    {'Measure':'Issues after deduplication','AI':11,'Human':7},
    {'Measure':'Implementations with findings','AI':int((a>0).sum()),'Human':int((b>0).sum())},
    {'Measure':'Confirmed concerns after deduplication (researcher review)','AI':11,'Human':7},
    {'Measure':'Context-dependent Vulnerability issues','AI':0,'Human':0},
    {'Measure':'False-positive Vulnerability issues','AI':0,'Human':0},
    {'Measure':'Security Hotspots','AI':60,'Human':56},
    {'Measure':'Hotspots classified as concerns (researcher review)','AI':44,'Human':50},
    {'Measure':'Debug hotspots requiring deployment context','AI':16,'Human':6},
    {'Measure':'Mean raw findings per implementation','AI':float(a.mean()),'Human':float(b.mean())},
    {'Measure':'Median raw findings per implementation','AI':float(np.median(a)),'Human':float(np.median(b))},
    {'Measure':'Two-sided Mann–Whitney p-value (raw findings)','AI':float(u.pvalue),'Human':''}]
wb = openpyxl.Workbook(); wb.remove(wb.active)
prompt_rows = load(ROOT/'prompts/prompts.json')
for title, rows in [('Summary',summary),('Per-implementation',per_sample),('Vulnerabilities',vulnerabilities),('Hotspots',hotspots),('Sources',sources),('Prompts',prompt_rows)]:
    sheet = wb.create_sheet(title); sheet.append(list(rows[0]))
    for row in rows:
        sheet.append([json.dumps(v) if isinstance(v,(dict,list)) else v for v in row.values()])
    sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
    for row in sheet:
        for cell in row:
            cell.font=Font(name='Calibri',size=11,bold=cell.row==1)
            cell.alignment=Alignment(vertical='top',wrap_text=True)
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width=min(65,max(14,max(len(str(c.value or '')) for c in column)+2))
wb.save(ROOT/'reports/Code_Comparison_Results.xlsx')
paragraphs = [
    'Thirty AI implementations and thirty human implementations were compared using SonarQube Community 26.3.0.120487 with the recorded Sonar way profiles. The scan covered 30 AI files and 225 human files. The implementation, rather than the individual source file, is the comparison unit.',
    'AI code was generated on 24 March 2026 using ChatGPT through its web interface and GitHub Copilot Chat version 0.35.2 in Visual Studio Code Agent mode. Most recorded Copilot requests used Auto model selection; sample 28 used Gemini 3 Pro Preview. Multiple Copilot prompts shared a conversation. The researcher reported one run per prompt and no user-defined custom or memory instructions related to code generation. The ChatGPT model is not specified here because its historical label remains to be resolved.',
    'Six AI JavaScript files received syntax-only corrections on eleven lines to enable parsing. Both original and scanned versions are supplied, with the correction record. Human implementations were selected by core task; repository commits, selected source files and implementation differences are documented in data/human-sources.csv. Their sizes and development histories differ, so these results describe this sample rather than a causal effect of AI authorship.',
    'SonarQube reported 16 Vulnerability findings in AI code and 9 in human code, all classified BLOCKER by the tool. Deduplication yielded 11 AI and 7 human issues, affecting 11 AI and 6 human implementations. Overlapping reports of the same secret assignment were counted once.',
    'The researcher reported confirming all 18 deduplicated issues as security concerns through manual code review, with no false positives or context-dependent Vulnerability issues. Of 116 Security Hotspots, 94 were classified as concerns (44 AI, 50 human); 22 debug-mode hotspots (16 AI, 6 human) required deployment context. These decisions were transcribed from the researcher’s statements; no second-reviewer agreement or exploitation testing is reported.',
    f'The recorded comparison of raw finding counts was not statistically significant (two-sided Mann–Whitney p = {u.pvalue:.3f}). The observed difference does not establish that the groups have equal security. This p-value concerns raw findings, not deduplicated or manually classified issues.',
    'Source references are supplied for all thirty human implementations. Human code is not bundled while redistribution questions are being resolved. Licence findings are recorded in licences/Repository_Licences.csv.'
]
doc = Document(); doc.add_heading('Code security comparison',0)
for paragraph in paragraphs: doc.add_paragraph(paragraph)
table=doc.add_table(rows=1, cols=3);table.style='Table Grid'
for cell,title in zip(table.rows[0].cells,('Measure','AI','Human')):cell.text=title
for row in summary[:11]:
    for cell,value in zip(table.add_row().cells,row.values()):cell.text=str(value)
doc.save(ROOT/'reports/Code_Comparison_Summary.docx')
(ROOT/'reports/RESULTS.md').write_text('# Code security comparison\n\n'+'\n\n'.join(paragraphs)+'\n')
print('Verified 60 implementations, 25 raw findings, 18 deduplicated issues and 116 hotspots; statistics reproduce the archived results.')
